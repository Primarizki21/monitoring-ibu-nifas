from pathlib import Path
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def apply_page_setup(doc):
    for section in doc.sections:
        section.page_width = Inches(8.27)   # A4 width
        section.page_height = Inches(11.69) # A4 height
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    style_normal = doc.styles["Normal"]
    style_normal.font.name = "Arial"
    style_normal.font.size = Pt(12)  # Font size 12pt
    style_normal.font.color.rgb = RGBColor(30, 41, 59)
    style_normal.paragraph_format.line_spacing = 1.25
    style_normal.paragraph_format.space_after = Pt(8)
    style_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY  # Justified

def set_cell_background(cell, hex_color: str):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, val in [("w:top", top), ("w:bottom", bottom), ("w:left", left), ("w:right", right)]:
        node = OxmlElement(m)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)

def add_table_caption(doc, caption_text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(caption_text)
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(71, 85, 105)
    return p

def add_callout(doc, title: str, text=None, text_items=None, bg_color="F0FDF4", border_color="16A34A", title_color=(22, 101, 52)):
    content_input = text if text is not None else text_items
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.27)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

    p_title = cell.paragraphs[0]
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_title.paragraph_format.space_before = Pt(2)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run(title)
    r_title.bold = True
    r_title.font.name = "Arial"
    r_title.font.size = Pt(11)
    r_title.font.color.rgb = RGBColor(*title_color)

    if isinstance(content_input, str):
        lines = [line.strip() for line in content_input.split("\n") if line.strip()]
    elif content_input is not None:
        lines = content_input
    else:
        lines = []

    for item in lines:
        p_item = cell.add_paragraph()
        p_item.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_item.paragraph_format.space_before = Pt(2)
        p_item.paragraph_format.space_after = Pt(4)
        if isinstance(item, str):
            r_text = p_item.add_run(item)
            r_text.font.name = "Arial"
            r_text.font.size = Pt(10.5)
            r_text.font.color.rgb = RGBColor(30, 41, 59)
        elif isinstance(item, (list, tuple)):
            for run_def in item:
                if isinstance(run_def, str):
                    r = p_item.add_run(run_def)
                    r.font.name = "Arial"
                    r.font.size = Pt(10.5)
                    r.font.color.rgb = RGBColor(30, 41, 59)
                else:
                    text = run_def[0]
                    is_bold = run_def[1] if len(run_def) > 1 else False
                    is_italic = run_def[2] if len(run_def) > 2 else False
                    r = p_item.add_run(text)
                    r.bold = is_bold
                    r.italic = is_italic
                    r.font.name = "Arial"
                    r.font.size = Pt(10.5)
                    r.font.color.rgb = RGBColor(30, 41, 59)
    
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(4)
def format_table_headers(table, col_widths, hex_bg="0F4C81"):
    hdr_cells = table.rows[0].cells
    for i, cell in enumerate(hdr_cells):
        cell.width = col_widths[i]
        set_cell_background(cell, hex_bg)
        set_cell_margins(cell, top=140, bottom=140, left=160, right=160)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            for r in p.runs:
                r.bold = True
                r.font.name = "Arial"
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(255, 255, 255)

def add_heading_2(doc, text: str, page_break=True):
    if page_break:
        doc.add_page_break()
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(8)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = h.add_run(text)
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor(30, 58, 138)
    return h

def add_heading_3(doc, text: str):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(6)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = h.add_run(text)
    r.bold = True
    r.font.size = Pt(12.5)
    r.font.color.rgb = RGBColor(15, 23, 42)
    return h

def add_justified_paragraph(doc, text_runs):
    """
    text_runs: list of tuples (text, is_bold, is_italic) or string.
    """
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(8)
    if isinstance(text_runs, str):
        p.add_run(text_runs)
        return p
    for item in text_runs:
        if isinstance(item, str):
            p.add_run(item)
        else:
            text = item[0]
            is_bold = item[1] if len(item) > 1 else False
            is_italic = item[2] if len(item) > 2 else False
            r = p.add_run(text)
            r.bold = is_bold
            r.italic = is_italic
    return p
def create_numbering_id(doc):
    numbering_elm = doc.part.numbering_part._element
    existing_nums = numbering_elm.findall(qn('w:num'))
    max_id = max([int(n.get(qn('w:numId'))) for n in existing_nums], default=0)
    new_id = max_id + 1
    new_num = parse_xml(f'<w:num {nsdecls("w")} w:numId="{new_id}"><w:abstractNumId w:val="7"/></w:num>')
    numbering_elm.append(new_num)
    return new_id

def add_numbered_paragraph(doc, text_runs, num_id: int):
    p = doc.add_paragraph(style="List Paragraph")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(4)
    pPr = p._p.get_or_add_pPr()
    numPr = parse_xml(f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="{num_id}"/></w:numPr>')
    pPr.append(numPr)
    if isinstance(text_runs, str):
        p.add_run(text_runs)
        return p
    for item in text_runs:
        if isinstance(item, str):
            p.add_run(item)
        else:
            text = item[0]
            is_bold = item[1] if len(item) > 1 else False
            is_italic = item[2] if len(item) > 2 else False
            r = p.add_run(text)
            r.bold = is_bold
            r.italic = is_italic
    return p

def add_cover_page(doc, title: str, subtitle: str, meta_lines: list):
    """
    Creates a dedicated cover page where all content is centered on the page.
    """
    # Push content vertically into the center of page 1
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(140)
    p_spacer.paragraph_format.space_after = Pt(0)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(8)
    r_title = p_title.add_run(title)
    r_title.bold = True
    r_title.font.size = Pt(22)
    r_title.font.color.rgb = RGBColor(15, 76, 129)

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(28)
    r_sub = p_sub.add_run(subtitle)
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = RGBColor(71, 85, 105)

    # Metadata Box Centered
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(5.8)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=160, bottom=160, left=180, right=180)

    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="8" w:space="0" w:color="CBD5E1"/>'
        f'<w:left w:val="single" w:sz="8" w:space="0" w:color="CBD5E1"/>'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="CBD5E1"/>'
        f'<w:right w:val="single" w:sz="8" w:space="0" w:color="CBD5E1"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

    p_box = cell.paragraphs[0]
    p_box.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_box.paragraph_format.space_before = Pt(4)
    p_box.paragraph_format.space_after = Pt(4)
    r_boxtitle = p_box.add_run("INFORMASI PROYEK PENGABDIAN MASYARAKAT")
    r_boxtitle.bold = True
    r_boxtitle.font.size = Pt(11)
    r_boxtitle.font.color.rgb = RGBColor(15, 76, 129)

    for line in meta_lines:
        p_line = cell.add_paragraph()
        p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_line.paragraph_format.space_before = Pt(2)
        p_line.paragraph_format.space_after = Pt(2)
        r = p_line.add_run(line)
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(30, 41, 59)

    # End of Page 1 -> Page Break to Page 2
    doc.add_page_break()

def add_chart_figure(doc, image_path: Path, caption_title: str, explanation_text: str, width=Inches(4.2)):
    """
    Inserts a centered flowchart image, figure caption, and an explanatory paragraph below it.
    """
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(10)
    p_img.paragraph_format.space_after = Pt(4)
    run_img = p_img.add_run()
    run_img.add_picture(str(image_path), width=width)

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(8)
    r_cap = p_cap.add_run(caption_title)
    r_cap.bold = True
    r_cap.font.size = Pt(10)
    r_cap.font.color.rgb = RGBColor(71, 85, 105)

    p_exp = doc.add_paragraph()
    p_exp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_exp.paragraph_format.space_before = Pt(4)
    p_exp.paragraph_format.space_after = Pt(12)
    r_exp = p_exp.add_run(explanation_text)
    r_exp.font.size = Pt(11)
    r_exp.font.color.rgb = RGBColor(51, 65, 85)
