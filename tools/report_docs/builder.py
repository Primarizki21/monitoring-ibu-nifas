from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT

from tools.report_docs.styles import (
    apply_page_setup,
    add_cover_page,
    add_callout,
    format_table_headers,
    set_cell_margins,
    set_cell_background,
    add_heading_2,
    add_heading_3,
    add_justified_paragraph,
    add_chart_figure,
    add_table_caption,
    create_numbering_id,
    add_numbered_paragraph,
)
from tools.report_docs import content

def build_project_charter_doc(assets_dir: Path) -> Document:
    doc = Document()
    apply_page_setup(doc)

    # 1. Dedicated Cover Page (Centered on Page 1, Page-break to Page 2)
    add_cover_page(
        doc,
        title=content.METADATA["title"],
        subtitle=content.METADATA["subtitle"],
        meta_lines=content.METADATA["meta_lines"]
    )

    # 2. Bab 1. Ringkasan Eksekutif (Starts directly on Page 2)
    add_heading_2(doc, "Bab 1. Ringkasan Eksekutif (Executive Summary)", page_break=False)
    for p_runs in content.CHAPTER_1_TEXTS:
        add_justified_paragraph(doc, p_runs)

    # 3. Bab 2. Prinsip Desain & Keputusan Kritis Sistem
    add_heading_2(doc, "Bab 2. Prinsip Desain & Keputusan Kritis Sistem", page_break=True)
    for item in content.CHAPTER_2_ITEMS:
        sub_title = item[0]
        intro_runs = item[1]
        numbered_items = item[2]
        add_heading_3(doc, sub_title)
        if intro_runs:
            add_justified_paragraph(doc, intro_runs)
        if numbered_items:
            num_id = create_numbering_id(doc)
            for n_runs in numbered_items:
                add_numbered_paragraph(doc, n_runs, num_id=num_id)
    # 4. Bab 3. Diagram & Alur Kerja Sistem (Flowchart Operasional)
    add_heading_2(doc, "Bab 3. Diagram & Alur Kerja Sistem (Flowchart Operasional)", page_break=True)
    add_justified_paragraph(
        doc,
        "Alur operasional sistem pemantauan nifas dibagi menjadi tiga tahapan visual low-level, "
        "menghubungkan faskes penolong persalinan, sistem automasi server, interaksi WhatsApp ibu, hingga intervensi bidan wilayah:"
    )

    # Render each sub-flowchart section with image + caption + explanation
    for sec in content.FLOWCHART_SECTIONS:
        add_heading_3(doc, sec["title"])
        img_path = assets_dir / sec["image_file"]
        if not img_path.exists():
            raise FileNotFoundError(f"Flowchart image not found: {img_path}")
        add_chart_figure(
            doc,
            image_path=img_path,
            caption_title=sec["caption"],
            explanation_text=sec["explanation"],
            width=Inches(4.6)
        )

    # Summary 5-Stage Table
    add_heading_3(doc, "3.4. Matriks Ringkasan Alur 5 Tahap Operasional")
    add_justified_paragraph(doc, content.FLOW_TABLE_EXPLANATION)
    add_table_caption(doc, content.FLOW_TABLE_CAPTION)
    tbl_flow = doc.add_table(rows=len(content.FLOW_TABLE_DATA), cols=6)
    tbl_flow.alignment = WD_TABLE_ALIGNMENT.CENTER
    flow_widths = [Inches(0.45), Inches(1.05), Inches(1.15), Inches(1.20), Inches(1.22), Inches(1.20)]
    for r_idx, row in enumerate(tbl_flow.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.text = content.FLOW_TABLE_DATA[r_idx][c_idx]
            cell.width = flow_widths[c_idx]
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            if r_idx % 2 == 1 and r_idx > 0:
                set_cell_background(cell, "F8FAFC")
    format_table_headers(tbl_flow, flow_widths, "0F4C81")
    p_note = doc.add_paragraph()
    p_note.paragraph_format.space_before = Pt(6)
    p_note.paragraph_format.space_after = Pt(4)
    r_n = p_note.add_run("Catatan: ")
    r_n.bold = True
    p_note.add_run("Spesifikasi diagram alir ANSI/ISO lengkap dan kode Mermaid dapat diakses pada dokumen repositori: FLOWCHART.md.")

    # 5. Bab 4. Struktur Peran & Hak Akses (Multi-Tenant RBAC)
    add_heading_2(doc, "Bab 4. Struktur Peran & Hak Akses (Multi-Tenant RBAC)", page_break=True)
    add_justified_paragraph(
        doc,
        "Sistem menerapkan isolasi data faskes yang ketat untuk menjamin kerahasiaan dan privasi data medis pasien:"
    )
    add_justified_paragraph(doc, content.RBAC_TABLE_EXPLANATION)
    add_table_caption(doc, content.RBAC_TABLE_CAPTION)

    tbl_rbac = doc.add_table(rows=len(content.RBAC_TABLE_DATA), cols=4)
    tbl_rbac.alignment = WD_TABLE_ALIGNMENT.CENTER
    rbac_widths = [Inches(1.2), Inches(1.4), Inches(1.4), Inches(2.27)]
    for r_idx, row in enumerate(tbl_rbac.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.text = content.RBAC_TABLE_DATA[r_idx][c_idx]
            cell.width = rbac_widths[c_idx]
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
    format_table_headers(tbl_rbac, rbac_widths, "0F4C81")

    # 6. Bab 5. Spesifikasi Teknis & Infrastruktur (Penjelasan Konseptual)
    add_heading_2(doc, "Bab 5. Spesifikasi Teknis & Infrastruktur (Penjelasan Lengkap)", page_break=True)
    add_justified_paragraph(
        doc,
        "Untuk memudahkan pemahaman seluruh pemangku kepentingan lintas disiplin, berikut penjelasan teknis beserta analogi fungsinya:"
    )

    for title, tech, analogy in content.TECH_ITEMS:
        add_heading_3(doc, title)
        add_justified_paragraph(doc, tech)
        add_callout(
            doc,
            title="Penjelasan Konseptual & Analogi Sistem:",
            text=analogy,
            bg_color="F0FDF4",
            border_color="16A34A",
            title_color=(22, 101, 52)
        )

    # 7. Bab 6. Protokol Medis & Matriks Skrining (KF1 s/d KF4)
    add_heading_2(doc, "Bab 6. Protokol Medis & Matriks Skrining (KF1 s/d KF4)", page_break=True)
    add_justified_paragraph(
        doc,
        "Parameter skrining klinis diturunkan langsung dari standar Buku KIA Kementerian Kesehatan RI:"
    )
    add_justified_paragraph(doc, content.MEDICAL_TABLE_EXPLANATION)
    add_table_caption(doc, content.MEDICAL_TABLE_CAPTION)

    tbl_med = doc.add_table(rows=len(content.MEDICAL_TABLE_DATA), cols=4)
    tbl_med.alignment = WD_TABLE_ALIGNMENT.CENTER
    med_widths = [Inches(0.9), Inches(1.3), Inches(2.0), Inches(2.07)]
    for r_idx, row in enumerate(tbl_med.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.text = content.MEDICAL_TABLE_DATA[r_idx][c_idx]
            cell.width = med_widths[c_idx]
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
    format_table_headers(tbl_med, med_widths, "0F4C81")

    # 8. Bab 7. Rencana Anggaran Biaya & Efisiensi Finansial (Pilot 3 Bulan)
    add_heading_2(doc, "Bab 7. Rencana Anggaran Biaya & Efisiensi Finansial (Pilot 3 Bulan)", page_break=True)
    add_justified_paragraph(
        doc,
        "Rincian anggaran telah disinkronkan secara presisi dengan lembar kerja Google Spreadsheet:"
    )
    add_justified_paragraph(doc, content.BUDGET_TABLE_EXPLANATION)
    add_table_caption(doc, content.BUDGET_TABLE_CAPTION)

    tbl_rab = doc.add_table(rows=len(content.BUDGET_TABLE_DATA), cols=4)
    tbl_rab.alignment = WD_TABLE_ALIGNMENT.CENTER
    rab_widths = [Inches(2.1), Inches(2.0), Inches(1.1), Inches(1.07)]
    for r_idx, row in enumerate(tbl_rab.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.text = content.BUDGET_TABLE_DATA[r_idx][c_idx]
            cell.width = rab_widths[c_idx]
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            if r_idx == len(content.BUDGET_TABLE_DATA) - 1:
                set_cell_background(cell, "DCFCE7")
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.bold = True
                        r.font.color.rgb = RGBColor(22, 101, 52)

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(2)
    p_sp.paragraph_format.space_after = Pt(2)
    add_callout(
        doc,
        title="CATATAN EFISIENSI FINANSIAL",
        text=content.EFFICIENCY_NOTES,
        bg_color="FEF9C3",
        border_color="CA8A04",
        title_color=(133, 77, 14)
    )

    # 9. Bab 8. Roadmap 3 Bulan Live Deployment
    add_heading_2(doc, "Bab 8. Rencana Kerja Lapangan (Roadmap 3 Bulan Live Deployment)", page_break=True)
    add_justified_paragraph(
        doc,
        "Jadwal di bawah ini menegaskan tiga bulan masa operasional langsung di lapangan (Live Deployment) di mana sistem telah aktif melayani pasien nifas:"
    )
    add_justified_paragraph(doc, content.ROADMAP_TABLE_EXPLANATION)
    add_table_caption(doc, content.ROADMAP_TABLE_CAPTION)

    tbl_road = doc.add_table(rows=len(content.ROADMAP_TABLE_DATA), cols=3)
    tbl_road.alignment = WD_TABLE_ALIGNMENT.CENTER
    road_widths = [Inches(1.7), Inches(2.6), Inches(1.97)]
    for r_idx, row in enumerate(tbl_road.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.text = content.ROADMAP_TABLE_DATA[r_idx][c_idx]
            cell.width = road_widths[c_idx]
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
    format_table_headers(tbl_road, road_widths, "0F4C81")

    return doc
